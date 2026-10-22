"""Proofreading kit for the album's Armenian terminology (P1-10). Exports every album term —
sheet titles, room names, ceilings, electrical symbols, and the vocabulary shared with the UI —
to an XLSX the Armenian-speaking architect fills in (corrected HY + comment), then lists the
proposed changes for review before they are applied to `app/album/i18n.py` / the UI dictionaries.

    python -m app.album.glossary export album_terms.xlsx
    python -m app.album.glossary diff album_terms_reviewed.xlsx
"""
from __future__ import annotations

import sys
from pathlib import Path

from . import i18n

HEAD = ["table", "key", "hy", "ru", "en", "hy_corrected", "comment"]


def terms() -> list[dict]:
    rows = []
    for table, data in (("T", i18n.T), ("TX", i18n.TX), ("CEILING_TYPES", i18n.CEILING_TYPES), ("ROOM_NAMES", i18n.ROOM_NAMES)):
        rows += [{"table": table, "key": k, **{lang: v.get(lang, "") for lang in ("hy", "ru", "en")}} for k, v in data.items()]
    rows += [{"table": "ELEC", "key": k, "hy": v[0], "ru": v[1], "en": v[2]} for k, v in i18n.ELEC.items()]
    rows += [{"table": "shared", "key": k, **{lang: v.get(lang, "") for lang in ("hy", "ru", "en")}} for k, v in i18n._shared().items()]
    return rows


def export(path: Path) -> int:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "terms"
    ws.append(HEAD)
    for c in ws[1]:
        c.font = Font(bold=True)
    rows = terms()
    for r in rows:
        ws.append([r["table"], r["key"], r["hy"], r["ru"], r["en"], "", ""])
    for col, w in zip("ABCDEFG", (14, 28, 48, 48, 40, 48, 40), strict=True):
        ws.column_dimensions[col].width = w
    for row in ws.iter_rows(min_row=2, min_col=6, max_col=7):
        for c in row:
            c.fill = PatternFill("solid", fgColor="FFF6E6")
    ws.freeze_panes = "C2"
    wb.save(path)
    return len(rows)


def diff(path: Path) -> list[dict]:
    """Rows where the reviewer proposed a different Armenian term."""
    from openpyxl import load_workbook

    ws = load_workbook(path, read_only=True).active
    rows = list(ws.iter_rows(values_only=True))
    idx = {h: i for i, h in enumerate(rows[0])}
    out = []
    for r in rows[1:]:
        new = (r[idx["hy_corrected"]] or "").strip()
        if new and new != (r[idx["hy"]] or ""):
            out.append({"table": r[idx["table"]], "key": r[idx["key"]], "hy": r[idx["hy"]], "hy_corrected": new, "comment": r[idx["comment"]] or ""})
    return out


if __name__ == "__main__":
    cmd, p = sys.argv[1], Path(sys.argv[2])
    if cmd == "export":
        print(f"{export(p)} terms → {p}")
    else:
        for d in diff(p):
            print(f"{d['table']}.{d['key']}: {d['hy']!r} → {d['hy_corrected']!r}  {d['comment']}")
