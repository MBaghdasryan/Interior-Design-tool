"""Machine-readable exports: DXF (CAD, layer standard), XLSX/CSV (estimate &
specification), GLB (web 3D), JSON (native project). DWG is produced from the DXF
by Autodesk Platform Services Design Automation (see app/aps)."""
from __future__ import annotations

import csv
import io
import json
import math
import struct

import ezdxf
from ezdxf.enums import TextEntityAlignment
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from ..scene.geometry import along, footprint, polygon_area, wall_outline, wall_solid_segments
from ..scene.model import Scene

LAYERS = {
    "A-WALL": 7, "A-WALL-DEMO": 1, "A-WALL-NEW": 3, "A-DOOR": 4, "A-GLAZ": 5, "A-AREA": 8, "A-AREA-IDEN": 2, "I-FURN": 30,
    "E-POWR": 6, "E-LITE": 40, "P-PLMB": 150, "M-HVAC": 10, "A-DIMS": 9,
}


def to_dxf(scene: Scene, title: str = "") -> bytes:
    doc = ezdxf.new("R2018", setup=True)
    doc.units = ezdxf.units.MM
    for name, color in LAYERS.items():
        if name not in doc.layers:
            doc.layers.add(name, color=color)
    msp = doc.modelspace()

    def P(x, y):  # DXF is Y-up
        return (x, -y)

    for w in scene.walls:
        layer = {"existing": "A-WALL", "demolish": "A-WALL-DEMO", "new": "A-WALL-NEW"}[w.status]
        segs = wall_solid_segments(w, scene.openings) if w.status != "demolish" else [(0, math.hypot(w.b[0] - w.a[0], w.b[1] - w.a[1]))]
        for s, e in segs:
            msp.add_lwpolyline([P(*p) for p in wall_outline(w, s, e)], close=True, dxfattribs={"layer": layer})
        if w.status != "demolish":
            msp.add_line(P(*w.a), P(*w.b), dxfattribs={"layer": "A-DIMS", "linetype": "CENTER"} if "CENTER" in doc.linetypes else {"layer": "A-DIMS"})
    for o in scene.openings:
        w = scene.wall(o.wall_id)
        if not w:
            continue
        p0, p1 = along(w, o.offset), along(w, o.offset + o.width)
        layer = "A-GLAZ" if o.type == "window" else "A-DOOR"
        msp.add_line(P(*p0), P(*p1), dxfattribs={"layer": layer})
        if o.type in ("door", "entrance_door", "balcony_door"):
            hinge = p0 if o.hinge == "start" else p1
            ang = math.degrees(math.atan2(-(p1[1] - p0[1]), p1[0] - p0[0]))
            msp.add_arc(P(*hinge), o.width, ang if o.hinge == "start" else ang + 90, ang + 90 if o.hinge == "start" else ang + 180, dxfattribs={"layer": "A-DOOR"})
        msp.add_text(o.id, height=120, dxfattribs={"layer": layer}).set_placement(P((p0[0] + p1[0]) / 2, (p0[1] + p1[1]) / 2 - 250), align=TextEntityAlignment.MIDDLE_CENTER)
    for r in scene.rooms:
        msp.add_lwpolyline([P(*p) for p in r.polygon], close=True, dxfattribs={"layer": "A-AREA"})
        cx = sum(p[0] for p in r.polygon) / len(r.polygon)
        cy = sum(p[1] for p in r.polygon) / len(r.polygon)
        msp.add_text(f"{r.number}. {r.name}", height=160, dxfattribs={"layer": "A-AREA-IDEN"}).set_placement(P(cx, cy), align=TextEntityAlignment.MIDDLE_CENTER)
        msp.add_text(f"{polygon_area(r.polygon) / 1e6:.2f} m2", height=130, dxfattribs={"layer": "A-AREA-IDEN"}).set_placement(P(cx, cy + 250), align=TextEntityAlignment.MIDDLE_CENTER)
    for o in scene.objects:
        fp = footprint(o)
        msp.add_lwpolyline([P(*q) for q in fp.corners()], close=True, dxfattribs={"layer": "I-FURN"})
        msp.add_text(o.code, height=100, dxfattribs={"layer": "I-FURN"}).set_placement(P(o.x, o.y), align=TextEntityAlignment.MIDDLE_CENTER)
    for e in scene.electrical:
        msp.add_circle(P(*e.position), 60, dxfattribs={"layer": "E-POWR"})
        msp.add_text(f"{e.code} h{e.height:.0f}", height=60, dxfattribs={"layer": "E-POWR"}).set_placement(P(e.position[0] + 90, e.position[1]))
    for li in scene.lights:
        if li.end:
            msp.add_line(P(*li.position), P(*li.end), dxfattribs={"layer": "E-LITE"})
        else:
            msp.add_circle(P(*li.position), 80, dxfattribs={"layer": "E-LITE"})
    for f in scene.fixtures:
        if f.size:
            w, d = f.size
            msp.add_lwpolyline([P(f.position[0] - w / 2, f.position[1] - d / 2), P(f.position[0] + w / 2, f.position[1] - d / 2),
                                P(f.position[0] + w / 2, f.position[1] + d / 2), P(f.position[0] - w / 2, f.position[1] + d / 2)], close=True,
                               dxfattribs={"layer": "P-PLMB" if f.type.startswith(("riser", "water", "drain")) else "M-HVAC"})
    doc.header["$PROJECTNAME"] = title[:60]
    buf = io.StringIO()
    doc.write(buf)
    return buf.getvalue().encode("utf-8")


XL = {
    "sheet": {"hy": "Նախահաշիվ", "ru": "Смета", "en": "Estimate"},
    "generated": {"hy": "Ձևավորված է՝ {t} UTC · արժույթ AMD", "ru": "Сформирована: {t} UTC · валюта AMD", "en": "Generated {t} UTC · currency AMD"},
    "cols": {"hy": ["№", "Բաժին", "Կատեգորիա", "Անվանում", "Մատակարար", "Արտիկուլ", "Քանակ", "Միավոր", "Գին", "Գումար", "Գնի կարգավիճակ", "Ստուգված", "Հղում", "Նշում"],
             "ru": ["№", "Раздел", "Категория", "Наименование", "Поставщик", "SKU", "Кол-во", "Ед.", "Цена", "Сумма", "Статус цены", "Проверено", "Ссылка", "Примечание"],
             "en": ["#", "Section", "Category", "Name", "Supplier", "SKU", "Qty", "Unit", "Price", "Amount", "Price status", "Checked", "Link", "Note"]},
    "totals": {"hy": ["Ընդամենը (հայտնի գներ)", "Հաստատված", "Մոտավոր", "DEMO", "Միջակայք․ սկսած", "Միջակայք․ մինչև"],
               "ru": ["Итого (известные цены)", "Подтверждённые", "Ориентировочные", "DEMO", "Диапазон: от", "Диапазон: до"],
               "en": ["Total (known prices)", "Confirmed", "Estimated", "DEMO", "Range: from", "Range: to"]},
    "sections": {"hy": "Բաժիններ", "ru": "Разделы", "en": "Sections"},
    "sec_cols": {"hy": ["Բաժին", "Դիրքեր", "Գումար", "Բաժին %", "Ըստ հարցման"], "ru": ["Раздел", "Позиций", "Сумма", "Доля", "По запросу"], "en": ["Section", "Items", "Amount", "Share", "On request"]},
}


def estimate_xlsx(estimate: dict, title: str, lang: str = "ru") -> bytes:
    lang = lang if lang in ("hy", "ru", "en") else "ru"
    wb = Workbook()
    ws = wb.active
    ws.title = XL["sheet"][lang]
    head = Font(bold=True)
    fill = PatternFill("solid", fgColor="D9C9B7")
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([XL["generated"][lang].format(t=estimate.get("generated_at", "")[:16])])
    ws.append([])
    cols = XL["cols"][lang]
    ws.append(cols)
    for c in ws[4]:
        c.font, c.fill = head, fill
    for ln in estimate.get("lines", []):
        ws.append([ln.get("id"), ln.get("group"), ln.get("category"), ln.get("name"), ln.get("supplier"), ln.get("sku"), ln.get("qty"), ln.get("unit"),
                   ln.get("unit_price"), ln.get("total"), ln.get("price_status"), (ln.get("checked_at") or "")[:10], ln.get("url"), ln.get("note")])
    ws.append([])
    t = estimate.get("totals", {})
    for k, lab in zip(("known", "confirmed", "estimated", "demo", "range_low", "range_high"), XL["totals"][lang]):
        ws.append(["", lab, "", "", "", "", "", "", "", t.get(k)])
        ws.cell(ws.max_row, 2).font = head
    ws.append([])
    for d in estimate.get("disclaimers", []):
        ws.append(["", d])
    widths = [7, 12, 16, 46, 26, 22, 9, 8, 12, 14, 14, 12, 30, 40]
    for i, w in enumerate(widths):
        ws.column_dimensions[chr(65 + i)].width = w
    for row in ws.iter_rows(min_row=5):
        for c in row[8:10]:
            c.number_format = "#,##0"
            c.alignment = Alignment(horizontal="right")
    ws2 = wb.create_sheet(XL["sections"][lang])
    ws2.append(XL["sec_cols"][lang])
    for s in estimate.get("sections", []):
        ws2.append([s["label"], s["lines"], s["total"], s["share"], s["quote_required"]])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def estimate_csv(estimate: dict) -> bytes:
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    w.writerow(["id", "group", "category", "name", "supplier", "sku", "qty", "unit", "unit_price_amd", "total_amd", "price_status", "checked_at", "url"])
    for ln in estimate.get("lines", []):
        w.writerow([ln.get(k) for k in ("id", "group", "category", "name", "supplier", "sku", "qty", "unit", "unit_price", "total", "price_status", "checked_at", "url")])
    return ("﻿" + buf.getvalue()).encode("utf-8")


# ---------------------------------------------------------------- GLB
def _hex_rgb(h: str) -> list[float]:
    h = h.lstrip("#")
    try:
        return [int(h[i:i + 2], 16) / 255 for i in (0, 2, 4)] + [1.0]
    except ValueError:
        return [0.8, 0.8, 0.8, 1.0]


def _box(x0, y0, x1, y1, z0, z1):
    # glTF is Y-up, metres; plan X → X, plan Y → Z, height → Y
    s = 0.001
    v = [(x0, z0, y0), (x1, z0, y0), (x1, z0, y1), (x0, z0, y1), (x0, z1, y0), (x1, z1, y0), (x1, z1, y1), (x0, z1, y1)]
    verts = [(x * s, y * s, z * s) for x, y, z in v]
    idx = [0, 2, 1, 0, 3, 2, 4, 5, 6, 4, 6, 7, 0, 1, 5, 0, 5, 4, 1, 2, 6, 1, 6, 5, 2, 3, 7, 2, 7, 6, 3, 0, 4, 3, 4, 7]
    return verts, idx


def _prism(outline, z0, z1):
    s = 0.001
    n = len(outline)
    verts = [(x * s, z0 * s, y * s) for x, y in outline] + [(x * s, z1 * s, y * s) for x, y in outline]
    idx = []
    for i in range(n):
        j = (i + 1) % n
        idx += [i, j, n + j, i, n + j, n + i]
    idx += [n, n + 1, n + 2, n, n + 2, n + 3] if n == 4 else []
    return verts, idx


def to_glb(scene: Scene) -> bytes:
    meshes: list[tuple[str, list, list, str]] = []
    H = scene.ceiling_height
    for w in scene.walls:
        if w.status == "demolish":
            continue
        for s, e in wall_solid_segments(w, scene.openings):
            v, i = _prism(wall_outline(w, s, e), 0, H)
            meshes.append((f"wall_{w.id}", v, i, "#f1ece5"))
    for r in scene.rooms:
        # floor slab per room (bounding box; L-shaped rooms are approximated in the web preview)
        xs = [p[0] for p in r.polygon]
        ys = [p[1] for p in r.polygon]
        col = next((f.color_hex for f in scene.finishes if f.room_id == r.id and f.surface == "floor"), "#d8c7b0")
        v, i = _box(min(xs), min(ys), max(xs), max(ys), -20, 0)
        meshes.append((f"floor_{r.id}", v, i, col))
    for o in scene.objects:
        fp = footprint(o)
        v, i = _box(fp.x0, fp.y0, fp.x1, fp.y1, o.elevation, o.elevation + max(10, o.h))
        meshes.append((f"{o.code}_{o.category}", v, i, o.color_hex))
    bin_parts, buffer_views, accessors, gl_meshes, nodes, materials = [], [], [], [], [], []
    mat_index: dict[str, int] = {}
    offset = 0
    for name, verts, idx, color in meshes:
        if color not in mat_index:
            mat_index[color] = len(materials)
            materials.append({"pbrMetallicRoughness": {"baseColorFactor": _hex_rgb(color), "metallicFactor": 0.0, "roughnessFactor": 0.8}, "doubleSided": True})
        vb = b"".join(struct.pack("<3f", *p) for p in verts)
        ib = b"".join(struct.pack("<I", k) for k in idx)
        xs, ys, zs = zip(*verts)
        buffer_views.append({"buffer": 0, "byteOffset": offset, "byteLength": len(vb), "target": 34962})
        bin_parts.append(vb)
        offset += len(vb)
        buffer_views.append({"buffer": 0, "byteOffset": offset, "byteLength": len(ib), "target": 34963})
        bin_parts.append(ib)
        offset += len(ib)
        accessors.append({"bufferView": len(buffer_views) - 2, "componentType": 5126, "count": len(verts), "type": "VEC3",
                          "min": [min(xs), min(ys), min(zs)], "max": [max(xs), max(ys), max(zs)]})
        accessors.append({"bufferView": len(buffer_views) - 1, "componentType": 5125, "count": len(idx), "type": "SCALAR"})
        gl_meshes.append({"name": name, "primitives": [{"attributes": {"POSITION": len(accessors) - 2}, "indices": len(accessors) - 1, "material": mat_index[color]}]})
        nodes.append({"mesh": len(gl_meshes) - 1, "name": name})
    binary = b"".join(bin_parts)
    gltf = {"asset": {"version": "2.0", "generator": "Renovara scene exporter"}, "scene": 0, "scenes": [{"nodes": list(range(len(nodes)))}],
            "nodes": nodes, "meshes": gl_meshes, "materials": materials, "accessors": accessors, "bufferViews": buffer_views,
            "buffers": [{"byteLength": len(binary)}]}
    js = json.dumps(gltf, separators=(",", ":")).encode()
    js += b" " * ((4 - len(js) % 4) % 4)
    binary += b"\0" * ((4 - len(binary) % 4) % 4)
    total = 12 + 8 + len(js) + 8 + len(binary)
    return struct.pack("<4sII", b"glTF", 2, total) + struct.pack("<I4s", len(js), b"JSON") + js + struct.pack("<I4s", len(binary), b"BIN\0") + binary
