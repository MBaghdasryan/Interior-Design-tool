"""Artifact-level validation of produced files (not just "a file exists")."""
import csv
import io
import json
import struct

import ezdxf
import pytest
from openpyxl import load_workbook

from app.album.draw import AlbumMeta
from app.album.sheets import AlbumBuilder
from app.exports.files import estimate_csv, estimate_xlsx, to_dxf, to_glb
from app.scene.rules import run_rules
from tests.test_design_core import designed, pool  # noqa: F401  (fixtures)

A3 = (1190.55, 841.89)  # points, landscape


@pytest.fixture(scope="module")
def album(designed):  # noqa: F811
    base, res, _, est = designed
    meta = AlbumMeta(lang="hy", company="Studio", tagline="interior design", phone="+374", address="Երևան", email="s@x", project_address="Demo",
                     client="", designer="Designer", accuracy_label="Verified Design", demo=True)
    pdf, sheets = AlbumBuilder(res.scene, base, est, run_rules(res.scene, {}), meta, project={"confidence": 0.9}, variant={"title": "V", "style": "warm_minimal"},
                               viewer_url="http://x/client/t", approvals=[]).build()
    return pdf, sheets


def test_pdf_a3_pages_text_and_fonts(album):
    pymupdf = pytest.importorskip("pymupdf")
    pdf, sheets = album
    doc = pymupdf.open(stream=pdf, filetype="pdf")
    assert doc.page_count == len(sheets) + 2  # cover + contents + sheets
    for page in doc:
        w, h = page.rect.width, page.rect.height
        assert abs(w - A3[0]) < 1 and abs(h - A3[1]) < 1, "every sheet is A3 landscape"
    # title block numbering like the reference album: «ԴՆ1.1» and «/N»
    p3 = doc[2].get_text()
    assert "ԴՆ1.1" in p3 and f"/{doc.page_count}" in p3
    assert "Չափագրություն" in p3  # Armenian sheet title rendered as text (not outlines)
    fonts = {f[3] for page in doc for f in page.get_fonts()}
    assert any("DejaVu" in f for f in fonts), fonts  # Unicode font embedded for Armenian/Cyrillic
    assert all(f[1] != "n/a" for page in doc for f in page.get_fonts())  # embedded, not referenced
    # every section is present and every sheet carries its own code
    codes = [s["code"] for s in sheets]
    assert len(codes) == len(set(codes))
    for i, s in enumerate(sheets):
        assert s["code"] in doc[i + 2].get_text()


def test_dxf_reads_back_and_passes_audit(designed):  # noqa: F811
    _, res, _, _ = designed
    data = to_dxf(res.scene, "demo")
    doc = ezdxf.read(io.StringIO(data.decode()))
    auditor = doc.audit()
    assert not auditor.has_errors, [str(e) for e in auditor.errors][:5]
    layers = {e.dxf.layer for e in doc.modelspace()}
    assert {"A-WALL", "A-WALL-NEW", "A-DOOR", "A-GLAZ", "A-AREA", "I-FURN", "E-POWR", "E-LITE"} <= layers
    assert doc.units == ezdxf.units.MM


def test_glb_is_valid_gltf_with_scene_inventory(designed):  # noqa: F811
    _, res, _, _ = designed
    glb = to_glb(res.scene)
    magic, version, length = struct.unpack("<4sII", glb[:12])
    assert magic == b"glTF" and version == 2 and length == len(glb)
    jlen, jtype = struct.unpack("<I4s", glb[12:20])
    assert jtype == b"JSON"
    gltf = json.loads(glb[20:20 + jlen])
    blen, btype = struct.unpack("<I4s", glb[20 + jlen:28 + jlen])
    assert btype == b"BIN\0" and blen == gltf["buffers"][0]["byteLength"]
    names = [n["name"] for n in gltf["nodes"]]
    assert sum(n.startswith("wall_") for n in names) > 10
    assert sum(n.startswith("floor_") for n in names) == len(res.scene.rooms)
    assert len([n for n in names if not n.startswith(("wall_", "floor_"))]) == len(res.scene.objects)
    for acc in gltf["accessors"]:
        bv = gltf["bufferViews"][acc["bufferView"]]
        size = {"VEC3": 12, "SCALAR": 4}[acc["type"]]
        assert acc["count"] * size == bv["byteLength"]


def test_xlsx_values_match_estimate(designed):  # noqa: F811
    _, _, _, est = designed
    for lang, header in (("hy", "Անվանում"), ("ru", "Наименование"), ("en", "Name")):
        wb = load_workbook(io.BytesIO(estimate_xlsx(est, "Demo", lang)))
        ws = wb.worksheets[0]
        assert ws.cell(4, 4).value == header
        totals = [r for r in ws.iter_rows(values_only=True) if r[0] is None and r[1] and r[9] is not None]
        assert totals[0][9] == est["totals"]["known"]
        line_sum = sum(r[9] for r in ws.iter_rows(min_row=5, values_only=True) if isinstance(r[0], str) and r[0].startswith("L") and r[9] is not None)
        assert line_sum == est["totals"]["known"]


def test_csv_encoding_and_delimiter(designed):  # noqa: F811
    _, _, _, est = designed
    raw = estimate_csv(est)
    assert raw.startswith(b"\xef\xbb\xbf")  # UTF-8 BOM for Excel
    rows = list(csv.reader(io.StringIO(raw.decode("utf-8-sig")), delimiter=";"))
    assert rows[0][:3] == ["id", "group", "category"] and len(rows) == len(est["lines"]) + 1


def test_project_json_schema(designed):  # noqa: F811
    from app.scene.model import Scene

    _, res, _, est = designed
    payload = json.loads(json.dumps({"format": "forma-project", "version": 1, "scene": res.scene.model_dump(mode="json"), "estimate": est}))
    assert Scene.model_validate(payload["scene"]).objects  # round-trips through the scene schema
    assert payload["estimate"]["totals"]["known"] == est["totals"]["known"]


@pytest.mark.parametrize("lang", ["hy", "en"])
def test_album_has_no_russian_ui_text(designed, lang):  # noqa: F811
    """Every Cyrillic word in a hy/en album must come from project data (names, notes, supplier titles), never from sheet vocabulary."""
    import re

    pymupdf = pytest.importorskip("pymupdf")
    base, res, _, est = designed
    meta = AlbumMeta(lang=lang, company="Studio", tagline="interior design", phone="", address="", email="", project_address="Demo",
                     client="", designer="Designer", accuracy_label="Concept Estimate", demo=True)
    issues = run_rules(res.scene, {})
    pdf, _ = AlbumBuilder(res.scene, base, est, issues, meta, project={"confidence": 0.9}, variant={"title": "Тест", "style": "warm_minimal", "tier": "comfort"},
                          viewer_url="http://x/client/t", approvals=[]).build()
    # data corpus = what a person or a supplier typed: names, notes, product titles, supplier names — not system text
    scenes = [sc.model_dump(mode="json") for sc in (res.scene, base)]
    for sc in scenes:
        sc["electrical"] = [e for e in sc["electrical"] if not e.get("note_key")]
    products = [(ln["name"], ln.get("supplier")) for ln in est["lines"] if ln["group"] not in ("works", "delivery", "installation", "contingency")]
    corpus = json.dumps([scenes, products], ensure_ascii=False)
    words = {w for page in pymupdf.open(stream=pdf, filetype="pdf") for w in re.findall(r"[А-Яа-яЁё]{2,}", page.get_text())}
    leaked = sorted(w for w in words if w not in corpus)
    assert not leaked, leaked[:30]


def test_album_glossary_export_and_review_diff(tmp_path):
    """P1-10 proofreading kit: every album term goes to the reviewer; corrections come back as a diff."""
    from openpyxl import load_workbook

    from app.album import glossary

    p = tmp_path / "terms.xlsx"
    n = glossary.export(p)
    assert n > 150
    wb = load_workbook(p)
    ws = wb.active
    assert [c.value for c in ws[1]] == glossary.HEAD
    keys = {(r[0].value, r[1].value): r for r in ws.iter_rows(min_row=2)}
    assert keys[("ROOM_NAMES", "bathroom")][2].value == "Սանհանգույց"
    keys[("ROOM_NAMES", "bathroom")][5].value = "Լոգասենյակ"
    keys[("ROOM_NAMES", "bathroom")][6].value = "test"
    wb.save(p)
    assert glossary.diff(p) == [{"table": "ROOM_NAMES", "key": "bathroom", "hy": "Սանհանգույց", "hy_corrected": "Լոգասենյակ", "comment": "test"}]


def test_glossary_review_overrides_album_terms(client):
    """Linguist proposes and approves an Armenian correction → the album vocabulary uses it; others cannot."""
    from app.album import i18n
    from app.db import SessionLocal
    from app.routers.glossary import approved_overrides, glossary_stamp
    from tests.conftest import login
    from tests.test_auth import STRONG, register

    register(client, "linguist@example.com")
    client.cookies.clear()
    admin = login(client, "admin@forma.local")
    assert client.post("/api/v1/admin/roles", headers=admin, json={"email": "linguist@example.com", "role": "linguist"}).json()["ok"]
    lh = {"Authorization": f"Bearer {client.post('/api/v1/auth/login', json={'email': 'linguist@example.com', 'password': STRONG}).json()['token']}"}
    client.cookies.clear()
    assert client.get("/api/v1/glossary", headers=login(client)).status_code == 403  # designer is not a reviewer
    r = client.get("/api/v1/glossary", headers=lh, params={"q": "bathroom"}).json()
    assert r["total"] > 300 and any(x["term_key"] == "ROOM_NAMES.bathroom" for x in r["items"])
    with SessionLocal() as db:
        before = glossary_stamp(db)
    e = client.put("/api/v1/glossary/ROOM_NAMES.bathroom", headers=lh, json={"hy": "Լոգասենյակ", "comment": "ճարտարապետական տերմին"}).json()
    assert e["status"] == "proposed"
    with SessionLocal() as db:
        assert "ROOM_NAMES.bathroom" not in approved_overrides(db)  # proposals are not applied
    assert client.post("/api/v1/glossary/ROOM_NAMES.bathroom/decision", headers=lh, json={"decision": "approved"}).json()["status"] == "approved"
    with SessionLocal() as db:
        i18n.set_overrides(approved_overrides(db))
        assert glossary_stamp(db) != before
    room = type("R", (), {"type": "bathroom", "name": "Санузел"})()
    assert i18n.room_name(room, "hy") == "Լոգասենյակ" and i18n.room_name(room, "en") == "Bathroom"
    assert client.get("/api/v1/glossary/export.xlsx", headers=lh).content[:2] == b"PK"
    # revert to the original wording removes the entry
    assert client.put("/api/v1/glossary/ROOM_NAMES.bathroom", headers=lh, json={"hy": "Սանհանգույց"}).json()["status"] == "original"
    i18n.set_overrides({})
